#!/usr/bin/env python3
"""Convert ViT PIM xlsx files to unified CSV format compatible with unified_database.csv.

Reads ViT-B16.xlsx, ViT-L16.xlsx, ViT-H14.xlsx from timeloop_experiments/
and converts them to the same column format as unified_database.csv.

Changes from raw xlsx:
  1. Rename duplicate fused_layer_type column → arch_target (value: 'PIM')
  2. Normalize net names: ViT_B16_prefill → vit_b16_s197, etc.
  3. Map layer names to match database convention
  4. Expand layer0_q/k/v/attn_o_proj → layer0_q_proj (divide metrics by 4)
  5. Generate all 4 fused_layer_type variants (single/start/middle/end)
  6. Unit conversion: latency ms→s, power W → energy J
  7. Set sequence_length=1 (sequence baked into workload dims)
  8. Drop kv_cache_length column

Usage:
  python3 convert_vit_xlsx.py
  python3 convert_vit_xlsx.py --output vit_pim_database.csv
"""

import os
import sys
import argparse
import openpyxl
import csv

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))

# CSV column order (same as unified_database.csv)
CSV_COLUMNS = [
    'net', 'layer_name', 'batch_size', 'sequence_length', 'mapper_idx',
    'fused_layer_type', 'tp_degree', 'arch_target', 'glb_scale',
    'pe_x_scale', 'pe_y_scale', 'dram_i', 'dram_o', 'latency',
    'static_power', 'dynamic_energy', 'area', 'utilization',
    'i_access', 'w_access', 'o_access'
]

# Net name mapping: xlsx name → database name
VIT_NET_MAP = {
    'ViT_B16_prefill': 'vit_b16_s197',
    'ViT_L16_prefill': 'vit_l16_s197',
    'ViT_H14_prefill': 'vit_h14_s257',
}

# Layer name mapping: xlsx name → (list of database names, split_count)
# split_count: how many sub-layers this xlsx entry represents
# For softmax: PIM combines all 4 softmax steps into one entry;
# we split equally into the 4 individual components so they match
# the workload YAML files and VirtualNetwork loading.
VIT_LAYER_MAP = {
    'layer0_q/k/v/attn_o_proj': (['layer0_q_proj'], 4),   # 4 projections share same q_proj workload
    'layer1_qk':                (['layer0_attn_qk'], 1),
    'layer6_av':                (['layer0_attn_v'], 1),
    'layer8_ffn1':              (['layer0_gate_proj'], 1),
    'layer8_ffn2':              (['layer0_down_proj'], 1),
    'layer_sftmax*':            (['layer0_softmax_max', 'layer0_softmax_sub_exp',
                                  'layer0_softmax_sum', 'layer0_softmax_div'], 4),
}

FUSED_LAYER_TYPES = ['single', 'start', 'middle', 'end']

# xlsx column indices (since header has duplicate 'fused_layer_type')
COL_NET = 0
COL_LAYER = 1
COL_BATCH = 2
COL_SEQ = 3
COL_KV = 4
COL_MAPPER = 5
COL_FUSED = 6        # fused_layer_type ('single')
COL_TP = 7
COL_ARCH = 8         # fused_layer_type.1 (actually arch_target = 'PIM')
COL_GLB = 9
COL_PEX = 10
COL_PEY = 11
COL_DRAMI = 12
COL_DRAMO = 13
COL_LATENCY = 14     # in ms
COL_STATIC_POWER = 15  # in W
COL_DYNAMIC_E = 16   # actually power in W
COL_AREA = 17
COL_UTIL = 18


def convert_xlsx(xlsx_path):
    """Convert a single ViT PIM xlsx to list of row dicts."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(all_rows) < 2:
        print(f"  WARNING: {xlsx_path} has no data rows")
        return []

    results = []
    for raw in all_rows[1:]:
        if not raw or raw[COL_NET] is None:
            continue

        net_raw = str(raw[COL_NET]).strip()
        layer_raw = str(raw[COL_LAYER]).strip()

        # Map net name
        net = VIT_NET_MAP.get(net_raw)
        if net is None:
            print(f"  WARNING: unknown net '{net_raw}', skipping")
            continue

        # Map layer name
        mapping = VIT_LAYER_MAP.get(layer_raw)
        if mapping is None:
            print(f"  WARNING: unknown layer '{layer_raw}', skipping")
            continue
        csv_layers, split_count = mapping

        # Unit conversion
        latency_ms = float(raw[COL_LATENCY]) if raw[COL_LATENCY] is not None else 0
        latency_s = latency_ms / 1000.0

        power_w = float(raw[COL_DYNAMIC_E]) if raw[COL_DYNAMIC_E] is not None else 0
        energy_j = power_w * latency_s  # W × s = J

        static_power = float(raw[COL_STATIC_POWER]) if raw[COL_STATIC_POWER] is not None else 0
        area = float(raw[COL_AREA]) if raw[COL_AREA] is not None else ''
        utilization = float(raw[COL_UTIL]) if raw[COL_UTIL] is not None else 0

        tp_degree = int(float(raw[COL_TP])) if raw[COL_TP] is not None else 1

        # For split layers, divide latency and energy proportionally
        layer_latency = latency_s / split_count
        layer_energy = energy_j / split_count

        for csv_layer in csv_layers:
            # Generate all fused_layer_type variants
            for flt in FUSED_LAYER_TYPES:
                row = {
                    'net': net,
                    'layer_name': csv_layer,
                    'batch_size': 1,
                    'sequence_length': 1,  # sequence baked into workload dims
                    'mapper_idx': 0,
                    'fused_layer_type': flt,
                    'tp_degree': tp_degree,
                    'arch_target': 'PIM',
                    'glb_scale': 1,
                    'pe_x_scale': 1,
                    'pe_y_scale': 1,
                    'dram_i': 'GDDR7',   # PIM is GDDR-based
                    'dram_o': 'GDDR7',
                    'latency': layer_latency,
                    'static_power': static_power,
                    'dynamic_energy': layer_energy,
                    'area': area,
                    'utilization': utilization,
                    'i_access': '',
                    'w_access': '',
                    'o_access': '',
                }
                results.append(row)

    return results


def main():
    parser = argparse.ArgumentParser(
        description="Convert ViT PIM xlsx files to unified CSV format")
    parser.add_argument("--output", default=os.path.join(_THIS_DIR, "vit_pim_database.csv"),
                        help="Output CSV path (default: vit_pim_database.csv)")
    parser.add_argument("--xlsx-dir", default=_THIS_DIR,
                        help="Directory containing ViT xlsx files (default: script dir)")
    args = parser.parse_args()

    xlsx_files = sorted([
        os.path.join(args.xlsx_dir, f)
        for f in os.listdir(args.xlsx_dir)
        if f.startswith('ViT') and f.endswith('.xlsx')
    ])

    if not xlsx_files:
        print("No ViT xlsx files found in %s" % args.xlsx_dir)
        return

    print("=" * 70)
    print("VIT PIM XLSX → CSV CONVERTER")
    print("=" * 70)

    all_results = []
    for xlsx_path in xlsx_files:
        fname = os.path.basename(xlsx_path)
        print(f"\nConverting {fname}...")
        rows = convert_xlsx(xlsx_path)
        print(f"  {len(rows)} rows generated")
        all_results.extend(rows)

    if not all_results:
        print("\nNo results to write.")
        return

    # Write CSV
    with open(args.output, 'w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(all_results)

    print(f"\nWrote {len(all_results)} rows to {args.output}")

    # Summary
    nets = sorted(set(r['net'] for r in all_results))
    layers = sorted(set(r['layer_name'] for r in all_results))
    tps = sorted(set(r['tp_degree'] for r in all_results))
    flts = sorted(set(r['fused_layer_type'] for r in all_results))

    print(f"\nSummary:")
    print(f"  Networks:         {nets}")
    print(f"  Layers:           {layers}")
    print(f"  TP degrees:       {tps}")
    print(f"  Fused layer types: {flts}")
    print(f"  Total rows:       {len(all_results)}")

    for net in nets:
        net_rows = [r for r in all_results if r['net'] == net]
        net_layers = sorted(set(r['layer_name'] for r in net_rows))
        print(f"\n  {net}:")
        for layer in net_layers:
            lr = [r for r in net_rows if r['layer_name'] == layer and r['fused_layer_type'] == 'single']
            for r in lr:
                print(f"    {layer:<25s} tp={r['tp_degree']} lat={r['latency']:.6f}s "
                      f"energy={r['dynamic_energy']:.6e}J")


if __name__ == '__main__':
    main()
